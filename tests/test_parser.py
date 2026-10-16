from __future__ import annotations

from pathlib import Path

from docx import Document

from legal_rag.chunking import make_chunks
from legal_rag.parser import build_sections, extract_blocks, heading
from legal_rag.types import Record


def fixture_docx(tmp_path: Path) -> Path:
    path = tmp_path / "synthetic.docx"
    doc = Document()
    doc.add_paragraph("ARTICLE XIII — MISCELLANEOUS")
    doc.add_paragraph("Section 13.9 — Survival")
    doc.add_paragraph("These provisions survive termination.")
    doc.add_paragraph("[Signature Page Follows]")
    doc.add_paragraph("IN WITNESS WHEREOF, the parties execute this Agreement.")
    doc.add_paragraph("ALDERSGATE SOFTWARE SOLUTIONS, INC.")
    doc.add_paragraph("EXHIBIT A")
    doc.add_paragraph("SERVICE DESCRIPTION")
    doc.add_paragraph("Schedule 4.17 sets forth all Material Contracts.")
    doc.add_paragraph("EXHIBIT B — SERVICE LEVELS")
    doc.add_paragraph("Section 2 — Fulfillment Services")
    doc.add_paragraph("The service will be provided.")
    doc.save(path)
    return path


def test_special_headings() -> None:
    assert heading("EXHIBIT A")["section_kind"] == "exhibit"
    assert heading("EXHIBIT B — SERVICE LEVELS")["section_kind"] == "exhibit"
    assert heading("[SIGNATURE PAGE FOLLOWS]")["section_kind"] == "signature_notice"
    assert heading("IN WITNESS WHEREOF, the parties execute this agreement.")["section_kind"] == "signature"
    assert heading("Schedule 4.17 sets forth a complete list") is None
    assert heading("This refers to Exhibit A") is None


def test_no_text_loss_and_exhibit_boundary(tmp_path: Path) -> None:
    path = fixture_docx(tmp_path)
    blocks = extract_blocks(path)
    sections = build_sections(path.name, blocks)
    chunks = make_chunks(sections, max_words=25)
    survival = next(s for s in sections if s["section_id"] == "13.9")
    assert "EXHIBIT A" not in survival["text"]
    assert "IN WITNESS WHEREOF" not in survival["text"]
    assert any(s["section_kind"] == "signature" for s in sections)
    assert any(s["section_kind"] == "exhibit" and s["section_id"] == "A" for s in sections)
    assert any(s["section_kind"] == "exhibit" and s["section_id"] == "B" for s in sections)
    assert any("Schedule 4.17 sets forth" in s["text"] for s in sections)
    assert any(c["section_id"] == "2" for c in chunks)
    assert "\n".join(s["text"] for s in sections) == "\n".join(b["text"] for b in blocks)


def test_appendix_region_and_bold_conjunction(tmp_path: Path) -> None:
    assert heading("and", bold_ratio=1.0) is None
    path = fixture_docx(tmp_path)
    sections = build_sections(path.name, extract_blocks(path))
    appendix_section = next(s for s in sections if s["section_id"] == "2")
    assert appendix_section["legal_ref"] == "Exhibit B / 2"
    assert "Exhibit B" in appendix_section["parent_path"]
    assert next(s for s in sections if s["section_id"] == "13.9")["document_region"] == "Main Agreement"


def test_token_split_lossless() -> None:
    from legal_rag.chunking import split_section_tokens
    class Tokenizer:
        def __call__(self, text: str, truncation: bool = False) -> Record:
            return {"input_ids": list(range(len(text.split()) + 2))}
    section = {"source": "x.docx", "text": "Section 1 — Sample\n" + ("agreement " * 100).rstrip(),
               "parent_path": [], "section_id": "1", "legal_ref": "1"}
    parts = split_section_tokens(section, Tokenizer(), max_tokens=25)
    assert len(parts) > 1
    assert "".join(parts) == section["text"]


def test_lettered_subclauses_preserve_parent_reference() -> None:
    blocks = [
        {"text": "Section 7.09 — Change of Control", "bold_ratio": 0, "style": "", "kind": "paragraph"},
        {"text": "(a) Change of Control as Event of Default", "bold_ratio": 0, "style": "", "kind": "paragraph"},
        {"text": "A change of control constitutes an event of default.", "bold_ratio": 0, "style": "", "kind": "paragraph"},
        {"text": "(b) Consequences of Change of Control Event of Default", "bold_ratio": 0, "style": "", "kind": "paragraph"},
        {"text": "The lender may accelerate repayment.", "bold_ratio": 0, "style": "", "kind": "paragraph"},
    ]
    sections = build_sections("credit.docx", blocks)
    assert [x["legal_ref"] for x in sections] == ["7.09", "7.09(a)", "7.09(b)"]
    assert "7.09 Change of Control" in sections[2]["parent_path"]
    assert heading("by and between", bold_ratio=1.0) is None
